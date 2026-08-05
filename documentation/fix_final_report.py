"""
In-place formatting/content fixes for final_report.docx, applied on top of the
user's manual edits (do NOT regenerate from generate_final_report.py, which
would wipe those edits).

Fixes applied:
  1. Re-bake Heading 1 / Heading 2 style colors (Word reset the run-level
     color to black on some headings after the last manual edit) and
     reapply explicit run formatting so every heading renders consistently.
  2. Restore missing decimal-numbering prefixes on 5 sub-headings
     (2.2, 2.3, 2.4, 2.5, 5.2) so the numbering scheme is parallel again.
  3. Clean up two sentences left grammatically broken by manual edits
     (Introduction scope sentence; the new 6.2 "role-specialized agents"
     bullet).
  4. Add numbered captions to the two main-text tables and the Appendix I
     figure, per guideline 1.2(vii): "illustrations ... should always be
     accompanied by a number and an appropriate title."
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

PATH = os.path.join(os.path.dirname(__file__), 'final_report.docx')
doc = Document(PATH)

BLUE_DARK = RGBColor(0x1F, 0x4E, 0x79)
BLUE_MID  = RGBColor(0x2E, 0x74, 0xB5)

# ── 1. Style-level bake + run-level reassert for all headings ──────────────────
def strip_theme_color(rpr):
    color_el = rpr.find(qn('w:color'))
    if color_el is not None and qn('w:themeColor') in color_el.attrib:
        del color_el.attrib[qn('w:themeColor')]
    shade_el = None
    if color_el is not None and qn('w:themeShade') in color_el.attrib:
        del color_el.attrib[qn('w:themeShade')]

for style_name, size, color in [('Heading 1', 14, BLUE_DARK), ('Heading 2', 12, BLUE_MID)]:
    st = doc.styles[style_name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    rpr = st.element.get_or_add_rPr()
    strip_theme_color(rpr)

for p in doc.paragraphs:
    if p.style.name == 'Heading 1':
        for r in p.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = BLUE_DARK
    elif p.style.name == 'Heading 2':
        for r in p.runs:
            r.font.name = 'Calibri'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = BLUE_MID

# ── 2. Restore missing decimal-numbering prefixes ───────────────────────────────
NUMBER_FIXES = {
    ' File Upload & EDA\n': '2.2  ',
    ' Query Execution — Phase 1: Planning\n': '2.3  ',
    ' Query Execution — Phase 2: Execution\n': '2.4  ',
    ' Session Memory Update\n': '2.5  ',
    ' Functional Test Outcomes': '5.2  ',
}
fixed_count = 0
for p in doc.paragraphs:
    if p.style.name == 'Heading 2' and p.text in NUMBER_FIXES and p.runs:
        p.runs[0].text = NUMBER_FIXES[p.text]
        fixed_count += 1
print(f'Restored numbering prefixes on {fixed_count} headings')

# ── 3. Clean up grammatically broken sentences ──────────────────────────────────
for p in doc.paragraphs:
    if p.text.startswith('The scope of this dissertation') and 'more analytic agents' in p.text:
        # collapse all runs into the first, carrying its formatting
        base_run = p.runs[0]
        for extra in list(p.runs[1:]):
            extra._r.getparent().remove(extra._r)
        base_run.text = (
            'The scope of this dissertation is limited to designing a conversational, '
            'dialogue-based data analysis agent that can write and execute code to '
            'analyze CSV and Parquet data files and summarize the results based on '
            'user requests, deployed as a locally hosted web application. It does not '
            'cover production-scale multi-tenant deployment, authentication, '
            'containerized execution sandboxing, or the addition of specialized '
            'multi-role agents, all of which are identified as future work in Section 6.'
        )
        print('Fixed Introduction scope sentence')

    if p.text.startswith('Agents for separate') or 'Agents for separate role' in p.text:
        runs = list(p.runs)
        # first 3 runs are the bold prefix ("Agents for separate" + "role" + ": ")
        for extra in runs[1:]:
            extra._r.getparent().remove(extra._r)
        runs[0].text = 'Role-specialized agents'
        runs[0].font.bold = True
        new_run = p.add_run(': ')
        new_run.font.bold = True
        new_run.font.name = 'Calibri'
        new_run.font.size = Pt(11)
        body_run = p.add_run(
            'Introduce additional, purpose-specialized agents for distinct classes of '
            'analytical tasks (e.g. forecasting, visualization, anomaly detection) so '
            'that each task type is routed to an agent built for it, rather than '
            'relying solely on the general-purpose Coder agent.'
        )
        body_run.font.bold = False
        body_run.font.name = 'Calibri'
        body_run.font.size = Pt(11)
        p.paragraph_format.line_spacing = 2.0
        print('Fixed 6.2 role-specialized-agents bullet')

# ── 4. Add numbered captions to tables and the architecture figure ─────────────
def make_caption_paragraph(text):
    cap = doc.add_paragraph()
    r = cap.add_run(text)
    r.font.bold = True
    r.font.italic = True
    r.font.size = Pt(10)
    r.font.name = 'Calibri'
    r.font.color.rgb = BLUE_DARK
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(4)
    cap.paragraph_format.line_spacing = 1.0
    return cap

# Table 1 caption: before the Agent Communication Flow table (identified by its header row)
for t in doc.tables:
    header_texts = [c.text for c in t.rows[0].cells]
    if header_texts == ['Source Agent', 'Target Agent', 'Signal / Content']:
        cap = make_caption_paragraph('Table 1: Inter-Agent Communication Signals')
        t._tbl.addprevious(cap._p)
        print('Added caption: Table 1')
    elif header_texts == ['#', 'Technical Parameter', 'Specification']:
        cap = make_caption_paragraph('Table 2: Major Technical Specifications')
        t._tbl.addprevious(cap._p)
        print('Added caption: Table 2')

# Figure 1 caption: after the architecture diagram picture in Appendix I
for p in doc.paragraphs:
    if p._p.xml.find('<pic:pic') != -1 or p._p.findall('.//' + qn('a:blip')):
        # this paragraph contains the inline image
        cap = make_caption_paragraph('Figure 1: System Architecture Diagram')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p._p.addnext(cap._p)
        print('Added caption: Figure 1')
        break

doc.save(PATH)
print(f'Saved: {PATH}')
