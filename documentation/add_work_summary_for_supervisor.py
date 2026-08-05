"""
Creates final_mentor_certificate_summary.docx: a new copy of
final_mentor_certificate.docx with an added, clearly-labeled "Summary of Work
Completed" page — factual background the Supervisor/Additional Examiner can
reference while writing their OWN evaluation. It does not fill in, suggest,
or imply any Excellent/Good/Fair/Poor rating, grade, or remark — those remain
entirely for the Supervisor and Additional Examiner to complete themselves.
"""
import shutil
import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = os.path.join(os.path.dirname(__file__), 'final_mentor_certificate.docx')
DST = os.path.join(os.path.dirname(__file__), 'final_mentor_certificate_summary.docx')

shutil.copy(SRC, DST)

BLUE_DARK = RGBColor(0x1F, 0x4E, 0x79)

doc = Document(DST)

doc.add_page_break()

h = doc.add_heading('SUMMARY OF WORK COMPLETED', level=1)
r = h.runs[0] if h.runs else h.add_run(h.text)
r.font.name = 'Calibri'
r.font.bold = True
r.font.size = Pt(14)
r.font.color.rgb = BLUE_DARK

note = doc.add_paragraph()
nr = note.add_run(
    'Prepared by the student, for the Supervisor’s and Additional Examiner’s '
    'reference only. This page is factual background on work completed — it is not an '
    'evaluation, rating, or remark, and carries no bearing on the Excellent / Good / Fair / '
    'Poor assessment above, which remains solely the Supervisor’s and Additional '
    'Examiner’s own judgment.'
)
nr.font.italic = True
nr.font.size = Pt(10)
nr.font.name = 'Calibri'
note.paragraph_format.space_after = Pt(12)

def body(text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.bold = bold
    p.paragraph_format.space_after = Pt(8)
    return p

def bullet(text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(18)
    lead = p.add_run('•  ')
    lead.font.size = Pt(11)
    lead.font.name = 'Calibri'
    if bold_prefix:
        br = p.add_run(bold_prefix + ': ')
        br.font.bold = True
        br.font.size = Pt(11)
        br.font.name = 'Calibri'
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.name = 'Calibri'
    p.paragraph_format.space_after = Pt(4)

def h2(text):
    p = doc.add_heading(text, level=2)
    r = p.runs[0] if p.runs else p.add_run(text)
    r.font.name = 'Calibri'
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    return p

h2('What was built')
body('AIE Autocode Agent is a two-phase multi-agent system that lets a business analyst '
     'query CSV/Parquet datasets in natural language, without writing code. In Phase 1 '
     '(Planning), a Manager and Metadata Specialist agent turn dataset EDA metadata and the '
     'user’s request into a structured Implementation Plan, reviewed and annotated by the '
     'user before approval. In Phase 2 (Execution), a Coder agent generates Python code, an '
     'Executor runs it, and a Feedback Agent validates the output, retrying until correctness '
     'is confirmed.')

h2('Testing carried out')
bullet('Functional testing across multiple real datasets, including a 504-row holiday-'
       'initiative dataset and a 276,577-row North America revenue dataset',
       'Datasets tested')
bullet('Tool-invocation mismatches and UI-streaming issues were identified and fixed across '
       'the build cycle (Builds 4.0-8.5)', 'Defects found & fixed')
bullet('Repeated runs completed within the Phase 2 retry limit and produced FeedbackAgent-'
       'approved, business-ready answers end-to-end', 'Outcome')

h2('Key results')
bullet('Both the holiday-initiative and revenue datasets produced correct, verified answers '
       'end-to-end', 'Functional correctness')
bullet('As an applied case study, the system computed WMAPE of 12.49% (USA) and 35.4% '
       '(Canada), correctly benchmarking the existing enterprise forecast as more accurate '
       'than the generated statistical baseline', 'Forecasting case study')

h2('Objectives met')
bullet('Built an intelligent multi-agent system for analyzing CSV/Parquet datasets')
bullet('Automated metadata extraction and implementation planning via analyzer agents')
bullet('Generated optimized analytical code via a dedicated Coder agent')
bullet('Validated code correctness through an automated feedback loop before presenting '
       'results to the analyst')

h2('Current status')
body('Dissertation Outline, Literature Review, Design & Development, and Testing are all '
     'complete; the final report and this evaluation form are being submitted for '
     'Supervisor and Additional Examiner review ahead of final submission.')

doc.save(DST)
print(f'Saved: {DST}')
