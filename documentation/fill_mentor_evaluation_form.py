"""
Fills the known administrative/identifying fields of
"Mentor Evaluation Form for Final Submission(1).docx" and saves the result as
a new file, final_mentor_certificate.docx, leaving the original untouched.

Deliberately NOT filled: the Excellent/Good/Fair/Poor tick-box ratings (Tables
0 and 1) and the "ENCIRCLE the Recommended Final Grade" line, and the
Signature / Place & Date rows in Table 2. Those are the Supervisor's and
Additional Examiner's own judgment/signature and must be completed by them —
filling them in on the student's behalf would misrepresent an evaluation that
hasn't actually happened.
"""
from docx import Document
from docx.shared import Pt
import os

SRC = os.path.join(os.path.dirname(__file__), 'Mentor Evaluation Form for Final Submission(1).docx')
DST = os.path.join(os.path.dirname(__file__), 'final_mentor_certificate.docx')

BITS_ID          = '2024DA04024'
STUDENT_NAME     = 'Hariharan B'
STUDENT_EMAIL    = '2024da04024@wilp.bits-pilani.ac.in'
SUPERVISOR_NAME  = 'Abhishek Hegde'
PROJECT_TITLE    = 'AIE AUTOCODE AGENT'

SUPERVISOR = {
    'Qualification': 'Post Graduate Diploma - PG, Data Science',
    'Designation': 'Lead Data Scientist',
    'Employing Organization & Location': 'O9 Solutions, Bangalore',
    'Mobile Number': '+91 9742417683',
    'Email Address': 'abhishek.hegde@o9solutions.com',
}
EXAMINER = {
    'Qualification': 'B.E Computer Science',
    'Designation': 'Senior Data Scientist',
    'Employing Organization & Location': 'O9 Solutions, Bangalore',
    'Mobile Number': '+91 7204592327',
    'Email Address': 'nitesh.mutkekar@o9solutions.com',
}

doc = Document(SRC)

def append_value(paragraph, value):
    r = paragraph.add_run(value)
    r.font.bold = False
    r.font.size = Pt(10)
    return r

FIELD_MATCHES = [
    ('BITS ID No.', BITS_ID),
    ('NAME OF THE STUDENT', STUDENT_NAME),
    ('EMAIL ADDRESS', STUDENT_EMAIL),
    ('NAME OF THE SUPERVISOR', SUPERVISOR_NAME),
    ('PROJECT TITLE', PROJECT_TITLE),
]
filled = set()
for p in doc.paragraphs:
    for label, value in FIELD_MATCHES:
        if p.text.strip().startswith(label) and label not in filled:
            append_value(p, value)
            filled.add(label)
            break
print('Filled header fields:', sorted(filled))

# Table 2: Particulars of Supervisor and Additional Examiner
particulars_table = None
for t in doc.tables:
    headers = [c.text.strip() for c in t.rows[0].cells]
    if headers == ['', 'Supervisor', 'Additional Examiner']:
        particulars_table = t
        break

if particulars_table is None:
    raise RuntimeError('Could not find the Supervisor/Additional Examiner particulars table')

ROW_KEY_MAP = {
    'Name': 'Name',
    'Qualification': 'Qualification',
    'Designation': 'Designation',
    'Employing Organization & Location': 'Employing Organization & Location',
    'Phone Number': None,          # left blank — only a mobile number is on record
    'Mobile Number': 'Mobile Number',
    'Email Address': 'Email Address',
    'Signature': None,             # requires an actual signature — left blank
    'Place & Date': None,          # filled at the time of signing — left blank
}
SUPERVISOR['Name'] = SUPERVISOR_NAME
EXAMINER['Name'] = 'Nitesh Mutkekar'

filled_rows = 0
for row in particulars_table.rows[1:]:
    row_label = row.cells[0].text.strip()
    key = ROW_KEY_MAP.get(row_label)
    if key is None:
        continue
    for cell, data in ((row.cells[1], SUPERVISOR), (row.cells[2], EXAMINER)):
        cell.text = data[key]
        for para in cell.paragraphs:
            for r in para.runs:
                r.font.size = Pt(10)
    filled_rows += 1
print(f'Filled {filled_rows} rows of the Supervisor/Additional Examiner particulars table')
print('Left blank (must be completed by the Supervisor/Additional Examiner themselves): '
      'evaluation tick-boxes, Recommended Final Grade, Signature, Place & Date')

doc.save(DST)
print(f'Saved: {DST}')
